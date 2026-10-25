"""Excel batch import of CRM client cards and their company cards."""

from __future__ import annotations

from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from crm.client_import import build_client_template_xlsx, export_clients_xlsx, import_clients_xlsx
from crm.models import Client, Company, CompanyMember

User = get_user_model()

HEADERS = ["Компания", "ИНН", "Юридический адрес", "Контакт", "Эл. почта", "Телефон", "Заметки"]


def _xlsx(rows: list[list[object]]) -> BytesIO:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    buf = BytesIO()
    book.save(buf)
    buf.seek(0)
    return buf


def _upload(rows: list[list[object]], name: str = "clients.xlsx") -> SimpleUploadedFile:
    payload = _xlsx(rows).getvalue()
    return SimpleUploadedFile(
        name,
        payload,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@pytest.mark.django_db
def test_template_download_is_importable_and_creates_nothing() -> None:
    """The shipped workbook has the parser headers; the example sheet is ignored."""
    payload = build_client_template_xlsx()
    book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    try:
        titles = book.sheetnames
        first = book[titles[0]]
        header = next(first.iter_rows(values_only=True))
    finally:
        book.close()
    assert titles[0] == "Клиенты"
    assert list(header) == HEADERS
    assert "Пример" in titles

    report = import_clients_xlsx(BytesIO(payload))
    assert report.created == 0
    assert report.bad_rows == 0
    assert Client.objects.count() == 0
    assert Company.objects.count() == 0


@pytest.mark.django_db
def test_import_creates_clients_and_one_company_card() -> None:
    """Two contacts of the same legal entity share one company; requisites are stored."""
    report = import_clients_xlsx(
        _xlsx(
            [
                HEADERS,
                [
                    "ООО «Ромашка»",
                    "7707083893",
                    "Москва, ул. Ленина, 1",
                    "Анна Иванова",
                    "Anna@Example.com",
                    "+7 495 111-11-11",
                    "закупка",
                ],
                [
                    'ООО "Ромашка"',
                    "",
                    "",
                    "Пётр Сидоров",
                    "petr@example.com",
                    "",
                    "",
                ],
            ],
        ),
    )
    assert report.created == 2
    assert report.companies_created == 1
    assert report.bad_rows == 0
    assert Client.objects.count() == 2
    anna = Client.objects.get(email="anna@example.com")
    petr = Client.objects.get(email="petr@example.com")
    assert anna.name == "Анна Иванова"
    assert anna.phone == "+7 495 111-11-11"
    assert anna.notes == "закупка"
    assert anna.company_ref_id == petr.company_ref_id
    org = anna.company_ref
    assert org is not None
    assert org.inn == "7707083893"
    assert org.legal_address == "Москва, ул. Ленина, 1"
    assert Company.objects.count() == 1
    assert CompanyMember.objects.filter(company=org, is_confirmed=True).count() == 2


@pytest.mark.django_db
def test_import_keeps_the_first_phone_when_the_cell_has_two() -> None:
    """A second number after «;» stays on the company and must not overflow the client key."""
    report = import_clients_xlsx(
        _xlsx(
            [
                HEADERS,
                [
                    "ООО Ромашка",
                    "",
                    "",
                    "Анна",
                    "anna-phone@example.com",
                    "+7 963 232-85-50; офис +7 927 244-63-31",
                    "",
                ],
            ],
        ),
    )
    assert report.created == 1
    assert report.bad_rows == 0
    client = Client.objects.get(email="anna-phone@example.com")
    assert client.phone_digits == "79632328550"
    org = client.company_ref
    assert org is not None
    assert "927 244-63-31" in org.phone


@pytest.mark.django_db
def test_import_swapped_phone_and_email_columns_still_creates_the_contact() -> None:
    """Платинум Групп: телефоны лежат в колонке почты, адрес — в колонке телефона."""
    report = import_clients_xlsx(
        _xlsx(
            [
                HEADERS,
                [
                    "ООО «Платинум Групп»",
                    "9718033383",
                    "",
                    "уточнить",
                    "+7 495 108-43-73; +7 916 221-69-96",
                    "—; info@agregat-group.ru",
                    "",
                ],
            ],
        ),
    )
    assert report.created == 1
    assert report.bad_rows == 0
    client = Client.objects.get(email="info@agregat-group.ru")
    assert client.phone_digits == "74951084373"
    org = client.company_ref
    assert org is not None
    assert org.inn == "9718033383"
    assert "info@agregat-group.ru" in org.email


@pytest.mark.django_db
def test_import_does_not_overwrite_existing_client_or_inn() -> None:
    """An email already in CRM stays untouched; a different INN does not replace the stored one."""
    Client.objects.create(name="Старое имя", email="anna@example.com", company="ООО Ромашка")
    org = Company.objects.get()
    org.inn = "7707083893"
    org.legal_address = "Старый адрес"
    org.save()

    report = import_clients_xlsx(
        _xlsx(
            [
                HEADERS,
                [
                    "ООО Ромашка",
                    "500100732259",
                    "Новый адрес",
                    "Новое имя",
                    "anna@example.com",
                    "+7 999 000-00-00",
                    "не записывать",
                ],
                [
                    "ООО Ромашка",
                    "500100732259",
                    "Новый адрес",
                    "Новый контакт",
                    "new@example.com",
                    "",
                    "",
                ],
            ],
        ),
    )
    assert report.created == 1
    assert report.skipped_existing == 1
    assert report.companies_created == 0
    anna = Client.objects.get(email="anna@example.com")
    assert anna.name == "Старое имя"
    assert anna.phone == ""
    assert anna.notes == ""
    org.refresh_from_db()
    assert org.inn == "7707083893"
    assert org.legal_address == "Старый адрес"
    assert any("не изменён" in item for item in report.warnings)
    created = Client.objects.get(email="new@example.com")
    assert created.company_ref_id == org.pk
    member = CompanyMember.objects.get(client=created)
    assert member.is_confirmed is True


@pytest.mark.django_db
def test_import_rejects_bad_rows_and_keeps_the_valid_one() -> None:
    """A broken email, a bad INN and a repeated address do not block the good row."""
    report = import_clients_xlsx(
        _xlsx(
            [
                HEADERS,
                ["ООО Ромашка", "", "", "Анна", "not-an-email", "", ""],
                ["ООО Ромашка", "7707083894", "", "Борис", "boris@example.com", "", ""],
                ["ООО Ромашка", "7707083893", "", "Анна", "anna@example.com", "", ""],
                ["ООО Ромашка", "", "", "Анна ещё раз", "anna@example.com", "", ""],
                ["", "", "", "", "", "", ""],
            ],
        ),
    )
    assert report.created == 1
    assert report.bad_rows == 3
    assert report.skipped_blank == 1
    assert Client.objects.filter(email="anna@example.com").exists()
    assert not Client.objects.filter(email="boris@example.com").exists()
    org = Company.objects.get()
    assert org.inn == "7707083893"


@pytest.mark.django_db
def test_import_rejects_workbook_without_required_headers() -> None:
    """A file that is not the template does not create cards."""
    from crm.client_import import ClientImportError

    with pytest.raises(ClientImportError):
        import_clients_xlsx(_xlsx([["foo", "bar"], ["a", "b"]]))
    assert Client.objects.count() == 0


@pytest.mark.django_db
def test_import_admin_requires_add_permission(client) -> None:
    """A signed-in staff user without add_client cannot upload or download the template."""
    staff = User.objects.create_user("crm-view", password="x", is_staff=True)
    view_perm = Permission.objects.get(content_type__app_label="crm", codename="view_client")
    staff.user_permissions.add(view_perm)
    client.force_login(staff)
    upload = _upload([HEADERS, ["ООО Ромашка", "", "", "Анна", "anna@example.com", "", ""]])
    response = client.post(reverse("admin:crm_client_import_clients"), {"file": upload})
    assert response.status_code == 403
    assert Client.objects.count() == 0
    template = client.get(reverse("admin:crm_client_import_clients_template"))
    assert template.status_code == 403
    listing = client.get(reverse("admin:crm_client_changelist"))
    assert listing.status_code == 200
    page = listing.content.decode()
    assert "Загрузить клиентов" not in page
    assert reverse("admin:crm_client_export_clients") in page


@pytest.mark.django_db
def test_import_admin_post_creates_cards_and_offers_template(client) -> None:
    """Superuser upload redirects to the client list and the template is a real xlsx."""
    admin = User.objects.create_superuser(username="crm-imp", email="imp@example.com", password="x")
    client.force_login(admin)
    page = client.get(reverse("admin:crm_client_import_clients"))
    assert page.status_code == 200
    html = page.content.decode()
    assert "Загрузить клиентов" in html
    assert 'enctype="multipart/form-data"' in html
    assert reverse("admin:crm_client_import_clients_template") in html

    upload = _upload(
        [
            HEADERS,
            [
                "ООО Ромашка",
                "7707083893",
                "Москва",
                "Анна",
                "anna@example.com",
                "",
                "",
            ],
        ],
    )
    response = client.post(reverse("admin:crm_client_import_clients"), {"file": upload})
    assert response.status_code == 302
    assert response.url == reverse("admin:crm_client_changelist")
    created = Client.objects.get(email="anna@example.com")
    assert created.assignee_id is None

    template = client.get(reverse("admin:crm_client_import_clients_template"))
    assert template.status_code == 200
    assert "spreadsheetml" in template["Content-Type"]
    assert "klienty-shablon.xlsx" in template["Content-Disposition"]
    assert template.content[:2] == b"PK"

    companies = client.get(reverse("admin:crm_company_changelist"))
    assert companies.status_code == 200
    company_html = companies.content.decode()
    assert reverse("admin:crm_client_import_clients") in company_html
    assert reverse("admin:crm_client_export_clients") in company_html


@pytest.mark.django_db
def test_manager_upload_keeps_the_new_client_on_their_list(client) -> None:
    """A manager only sees assigned cards, so the upload must assign the new ones."""
    from crm.services import scope_clients_for_manager

    manager = User.objects.create_user("mgr-imp", password="x", is_staff=True)
    view_perm = Permission.objects.get(content_type__app_label="crm", codename="view_client")
    add_perm = Permission.objects.get(content_type__app_label="crm", codename="add_client")
    manager.user_permissions.add(view_perm, add_perm)
    upload = _upload([HEADERS, ["ООО Ромашка", "", "", "Анна", "anna-mgr@example.com", "", ""]])
    client.force_login(manager)
    response = client.post(reverse("admin:crm_client_import_clients"), {"file": upload})
    assert response.status_code == 302
    created = Client.objects.get(email="anna-mgr@example.com")
    assert created.assignee_id == manager.pk
    assert created.company_ref is not None
    assert created.company_ref.assignee_id == manager.pk
    listing = client.get(reverse("admin:crm_client_changelist"))
    assert "anna-mgr@example.com" in listing.content.decode()
    stranger = User.objects.create_user("mgr-other-imp", password="x", is_staff=True)
    visible = scope_clients_for_manager(Client.objects.all(), stranger)
    assert visible.filter(pk=created.pk).count() == 0


@pytest.mark.django_db
def test_import_does_not_claim_a_company_whose_contacts_belong_to_someone_else(client) -> None:
    """A new row must not make the uploader the owner of another manager's company."""
    owner = User.objects.create_user("imp-owner", password="x", is_staff=True)
    uploader = User.objects.create_user("imp-uploader", password="x", is_staff=True)
    for codename in ("view_client", "add_client"):
        uploader.user_permissions.add(
            Permission.objects.get(content_type__app_label="crm", codename=codename),
        )
    Client.objects.create(
        name="Уже есть",
        email="kept@example.com",
        company="ООО Ромашка",
        assignee=owner,
    )
    upload = _upload([HEADERS, ["ООО Ромашка", "", "", "Новый", "new-row@example.com", "", ""]])
    client.force_login(uploader)
    response = client.post(reverse("admin:crm_client_import_clients"), {"file": upload})
    assert response.status_code == 302
    created = Client.objects.get(email="new-row@example.com")
    assert created.assignee_id == uploader.pk
    assert created.company_ref is not None
    assert created.company_ref.assignee_id is None
    kept = Client.objects.get(email="kept@example.com")
    assert kept.assignee_id == owner.pk


@pytest.mark.django_db
def test_export_matches_import_columns_and_does_not_duplicate() -> None:
    """The export is the import sheet: loading it again leaves existing cards as they are."""
    import_clients_xlsx(
        _xlsx(
            [
                HEADERS,
                [
                    "ООО «Ромашка»",
                    "7707083893",
                    "Москва, ул. Ленина, 1",
                    "Анна Иванова",
                    "anna@example.com",
                    "+7 495 111-11-11",
                    "закупка",
                ],
            ],
        ),
    )
    payload = export_clients_xlsx(Client.objects.all())
    book = load_workbook(BytesIO(payload), data_only=True)
    try:
        sheet = book.active
        assert sheet is not None
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    assert list(rows[0]) == HEADERS
    assert list(rows[1]) == [
        "ООО «Ромашка»",
        "7707083893",
        "Москва, ул. Ленина, 1",
        "Анна Иванова",
        "anna@example.com",
        "+7 495 111-11-11",
        "закупка",
    ]
    again = import_clients_xlsx(BytesIO(payload))
    assert again.created == 0
    assert again.skipped_existing == 1
    assert Client.objects.count() == 1
    assert Company.objects.get().inn == "7707083893"


@pytest.mark.django_db
def test_export_admin_hides_another_managers_clients(client) -> None:
    """A manager downloads only the cards the client list already shows them."""
    manager = User.objects.create_user("mgr-export", password="x", is_staff=True)
    other = User.objects.create_user("mgr-other", password="x", is_staff=True)
    view_perm = Permission.objects.get(content_type__app_label="crm", codename="view_client")
    manager.user_permissions.add(view_perm)
    Client.objects.create(
        name="Свой",
        email="mine@example.com",
        company="ООО Своя",
        assignee=manager,
    )
    Client.objects.create(
        name="Чужой",
        email="other@example.com",
        company="ООО Чужая",
        assignee=other,
    )
    stranger = User.objects.create_user("no-crm", password="x", is_staff=True)
    client.force_login(stranger)
    denied = client.get(reverse("admin:crm_client_export_clients"))
    assert denied.status_code == 403

    client.force_login(manager)
    response = client.get(reverse("admin:crm_client_export_clients"))
    assert response.status_code == 200
    assert "spreadsheetml" in response["Content-Type"]
    assert response["Content-Disposition"].startswith('attachment; filename="klienty-')
    book = load_workbook(BytesIO(response.content), data_only=True)
    try:
        sheet = book.active
        assert sheet is not None
        emails = [row[4] for row in sheet.iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()
    assert emails == ["mine@example.com"]


@pytest.mark.django_db
def test_lead_client_and_company_lists_show_a_memo(client) -> None:
    """Each daily list says what the page is for before the table."""
    from pathlib import Path

    admin = User.objects.create_superuser("memo-admin", "memo-admin@example.com", "password12")
    client.force_login(admin)
    leads = client.get(reverse("admin:leads_lead_changelist")).content.decode()
    people = client.get(reverse("admin:crm_client_changelist")).content.decode()
    firms = client.get(reverse("admin:crm_company_changelist")).content.decode()
    assert "Новые обращения с сайта" in leads
    assert "Карточка человека" in people
    assert "Карточка фирмы" in firms
    css = Path("static/admin/css/hoocon-unfold-extras.css").read_text(encoding="utf-8")
    dark = css.split(".dark .hoocon-page-memo {", 1)[1].split("}", 1)[0]
    assert "background:" in dark
    assert "#f5f5f4" in dark


@pytest.mark.django_db
def test_company_and_client_lists_keep_cards_or_rows(client) -> None:
    """Desktop lists remember cards or rows. An unknown view does not replace it."""
    admin = User.objects.create_superuser("list-view", "list-view@example.com", "password12")
    client.force_login(admin)
    companies = reverse("admin:crm_company_changelist")
    clients = reverse("admin:crm_client_changelist")

    page = client.get(companies)
    html = page.content.decode()
    assert 'data-hoocon-list-view="cards"' in html
    assert "Карточки" in html
    assert "Строки" in html
    assert "view=rows" in html

    chosen = client.get(f"{companies}?view=rows&q=атом")
    body = chosen.content.decode()
    assert 'data-hoocon-list-view="rows"' in body
    assert "view=cards" in body
    assert "q=%D0%B0%D1%82%D0%BE%D0%BC" in body or "q=атом" in body

    again = client.get(companies)
    assert 'data-hoocon-list-view="rows"' in again.content.decode()
    people = client.get(clients)
    assert 'data-hoocon-list-view="rows"' in people.content.decode()

    junk = client.get(f"{companies}?view=kanban")
    assert 'data-hoocon-list-view="rows"' in junk.content.decode()


@pytest.mark.django_db
def test_import_reads_clients_sheet_when_example_is_active() -> None:
    """An open «Пример» tab must not replace the «Клиенты» rows."""
    book = Workbook()
    sample = book.active
    assert sample is not None
    sample.title = "Пример"
    sample.append(HEADERS)
    sample.append(
        ["ООО «Пример»", "", "", "Иван Петров", "ivan.petrov@example.com", "", "образец"],
    )
    data = book.create_sheet("Клиенты", 0)
    data.append(HEADERS)
    data.append(
        ["ООО «Ромашка»", "", "", "Анна", "anna-sheet@example.com", "", "живая строка"],
    )
    book.active = sample
    buf = BytesIO()
    book.save(buf)
    buf.seek(0)
    loaded = load_workbook(buf, read_only=True, data_only=True)
    try:
        assert loaded.active is not None
        assert loaded.active.title == "Пример"
    finally:
        loaded.close()
    buf.seek(0)

    report = import_clients_xlsx(buf)

    assert report.created == 1
    assert report.bad_rows == 0
    assert Client.objects.filter(email="anna-sheet@example.com").exists()
    assert not Client.objects.filter(email="ivan.petrov@example.com").exists()


@pytest.mark.django_db
def test_duplicate_email_during_import_counts_as_already_present() -> None:
    """A unique email conflict is «already there», not a silent database error."""
    from django.db.models.signals import pre_save

    def steal(sender: object, instance: Client, **kwargs: object) -> None:
        if instance.pk or instance.email != "race@example.com":
            return
        pre_save.disconnect(steal, sender=Client)
        Client.objects.create(name="уже", email="race@example.com", company="Другая")

    pre_save.connect(steal, sender=Client)
    try:
        report = import_clients_xlsx(
            _xlsx([HEADERS, ["ООО «Гонка»", "", "", "Гонка", "race@example.com", "", ""]]),
        )
    finally:
        pre_save.disconnect(steal, sender=Client)

    assert report.skipped_existing == 1
    assert report.bad_rows == 0
    assert report.created == 0
    assert "race@example.com" not in " ".join(report.errors)


@pytest.mark.django_db
def test_non_email_integrity_error_is_a_row_error() -> None:
    """A dedupe-key conflict is reported. It is not counted as an existing client."""
    from django.db import IntegrityError
    from django.db.models.signals import pre_save

    def boom(sender: object, instance: Company, **kwargs: object) -> None:
        pre_save.disconnect(boom, sender=Company)
        raise IntegrityError(
            'duplicate key value violates unique constraint "crm_plannedtouch_dedupe_key_key"',
        )

    pre_save.connect(boom, sender=Company)
    try:
        report = import_clients_xlsx(
            _xlsx([HEADERS, ["ООО «Сигнал»", "", "", "Пётр", "petr-signal@example.com", "", ""]]),
        )
    finally:
        pre_save.disconnect(boom, sender=Company)

    assert report.created == 0
    assert report.skipped_existing == 0
    assert report.bad_rows == 1
    assert "не удалось записать" in report.errors[0]
    assert "petr-signal@example.com" not in report.errors[0]
    assert "dedupe_key" not in report.errors[0]
