"""Batch-create CRM client cards (and their company cards) from Excel.

One data row is one contact. The client card key is the email. Rows that
share a company name (ignoring case and quotes) attach to one Company;
INN and legal address are filled only while those fields are still empty.
"""

from __future__ import annotations

import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from typing import Any, BinaryIO

from django.core.exceptions import ValidationError
from django.core.validators import EmailValidator
from django.db import IntegrityError, transaction
from django.db.models import QuerySet
from django.db.utils import DataError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from crm.models import Client, Company, CompanyMember
from crm.outreach import CompanyProfile, apply_company_profile, ensure_work_labels, remember_client_plan
from crm.work_labels import (
    extract_contact_name,
    labels_for_note,
    loose_company_key,
    niche_label_code,
    niche_title,
    planned_call_date,
    region_label_code,
    relevance_label_code,
    relevance_value,
)
from leads.services import normalize_company_label

CLIENT_SHEET_HEADERS: tuple[str, ...] = (
    "Компания",
    "ИНН",
    "Юридический адрес",
    "Контакт",
    "Эл. почта",
    "Телефон",
    "Заметки",
)
_COLUMN_WIDTHS = (36, 16, 42, 28, 32, 20, 36)

MAX_CLIENT_IMPORT_ROWS = 1000
_MAX_NAME = 200
_MAX_COMPANY = 200
_MAX_ADDRESS = 400
_MAX_PHONE = 50
_MAX_NOTES = 5000

_EMAIL = EmailValidator()
_INN10_WEIGHTS = (2, 4, 10, 3, 5, 9, 4, 6, 8)
_INN12_WEIGHTS_11 = (7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
_INN12_WEIGHTS_12 = (3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8)
_HEADER_SPLIT_RE = re.compile(r"[\s._\-]+")

# Canonical header token → field. First match wins.
_HEADER_FIELDS: dict[str, str] = {
    "компания": "company",
    "название": "company",
    "организация": "company",
    "юрлицо": "company",
    "инн": "inn",
    "юридическийадрес": "legal_address",
    "юрадрес": "legal_address",
    "адрес": "legal_address",
    "контакт": "name",
    "имя": "name",
    "фио": "name",
    "элпочта": "email",
    "электроннаяпочта": "email",
    "email": "email",
    "почта": "email",
    "телефон": "phone",
    "тел": "phone",
    "заметки": "notes",
    "комментарий": "notes",
}
_BASE_HEADER_FIELDS: dict[str, str] = {
    "компания": "company",
    "город": "city",
    "сайт": "website",
    "телефон": "phone",
    "email": "email",
    "элпочта": "email",
    "профильизбазы": "profile",
    "профиль(избазы)": "profile",
    "профиль": "profile",
    "регионбазы": "region",
    "регион": "region",
    "релевантность": "relevance",
    "нишаhoocon": "niche",
    "ниша": "niche",
}
_PLACEHOLDER_CONTACTS = frozenset({"", "—", "-", "–", "уточнить"})
_REQUIRED_LABELS = {
    "company": "Компания",
    "name": "Контакт",
    "email": "Эл. почта",
}


class ClientImportError(ValueError):
    """Workbook cannot be imported (missing headers, unreadable, too large)."""


class _RowAbortedError(Exception):
    """This row must roll back; the message is already on the report."""


@dataclass
class _ClientRow:
    """One validated contact row, ready to insert."""

    row_number: int
    company: str
    inn: str
    legal_address: str
    name: str
    email: str
    phone: str
    notes: str
    emails: tuple[str, ...] = ()
    ogrn: str = ""
    website: str = ""
    city: str = ""
    niche: str = ""
    region: str = ""
    relevance: str = ""
    label_codes: tuple[str, ...] = ()
    next_contact_on: date | None = None


@dataclass
class ClientImportReport:
    """Counters and row messages from one client workbook."""

    created: int = 0
    companies_created: int = 0
    companies_reused: int = 0
    touches_created: int = 0
    skipped_existing: int = 0
    skipped_blank: int = 0
    bad_rows: int = 0
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def summary(self) -> str:
        """One message for the admin flash, with the first row problems."""
        text = (
            f"Создано карточек клиентов: {self.created}; "
            f"новых компаний: {self.companies_created}; "
            f"компаний обновлено: {self.companies_reused}; "
            f"плановых звонков: {self.touches_created}; "
            f"уже были: {self.skipped_existing}; "
            f"пустых строк: {self.skipped_blank}; "
            f"с ошибкой: {self.bad_rows}."
        )
        extra = self.errors[:8] + self.warnings[:5]
        if extra:
            text = f"{text} {' '.join(extra)}"
        return text


def header_token(value: Any) -> str:
    """Collapse a header cell to a comparable token."""
    text = " ".join(str(value or "").strip().casefold().split())
    text = text.replace("ё", "е")
    return _HEADER_SPLIT_RE.sub("", text)


def inn_checksum_ok(inn: str) -> bool:
    """Return True for a 10- or 12-digit INN with a valid checksum.

    Empty string is not valid. The caller treats a blank cell as "no INN".
    """
    if len(inn) == 10 and inn.isdigit():
        total = sum(int(digit) * weight for digit, weight in zip(inn[:9], _INN10_WEIGHTS, strict=True))
        return total % 11 % 10 == int(inn[9])
    if len(inn) == 12 and inn.isdigit():
        first = sum(int(digit) * weight for digit, weight in zip(inn[:10], _INN12_WEIGHTS_11, strict=True))
        if first % 11 % 10 != int(inn[10]):
            return False
        second = sum(int(digit) * weight for digit, weight in zip(inn[:11], _INN12_WEIGHTS_12, strict=True))
        return second % 11 % 10 == int(inn[11])
    return False


def _cell_text(value: Any) -> str:
    """Stringify a cell. Whole-number floats (Excel INN) drop the fraction."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _map_headers(header_row: tuple[Any, ...]) -> dict[str, int]:
    """Return field name → column index. Raises when a required column is missing."""
    found: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        field_name = _HEADER_FIELDS.get(header_token(cell))
        if field_name and field_name not in found:
            found[field_name] = idx
    missing = [label for key, label in _REQUIRED_LABELS.items() if key not in found]
    if missing:
        raise ClientImportError(
            "В первой строке листа «Клиенты» нужны колонки: "
            + ", ".join(f"«{label}»" for label in _REQUIRED_LABELS.values())
            + f". Не хватает: {', '.join(missing)}.",
        )
    return found


def _row_problem(row: _ClientRow) -> str | None:
    """Return a Russian reason to reject the row, or None when it can be inserted.

    Email is required only for a contact card. A row with a company and no
    address still creates the company card.
    """
    if not normalize_company_label(row.company):
        return "укажите компанию"
    if len(row.company) > _MAX_COMPANY:
        return "название компании длиннее 200 символов"
    if row.email and not row.name:
        return "укажите контакт"
    if len(row.name) > _MAX_NAME:
        return "имя контакта длиннее 200 символов"
    if row.email:
        if len(row.email) > 254:
            return "эл. почта длиннее 254 символов"
        try:
            _EMAIL(row.email)
        except ValidationError:
            return f"некорректная эл. почта «{row.email}»"
    if row.inn and not inn_checksum_ok(row.inn):
        return f"ИНН «{row.inn}» не проходит контрольную сумму"
    if len(row.legal_address) > _MAX_ADDRESS:
        return "юридический адрес длиннее 400 символов"
    if len(row.phone) > 300:
        return "телефон длиннее 300 символов"
    if len(row.notes) > _MAX_NOTES:
        return "заметки длиннее 5000 символов"
    return None


def _classify_inn(raw: str) -> tuple[str, str, str | None]:
    """Return ``(inn, ogrn, error)``. A 13-digit cell is an OGRN, not an INN."""
    digits = re.sub(r"\D", "", raw)
    if not digits:
        return "", "", None
    if len(digits) == 13:
        return "", digits, None
    if len(digits) in (10, 12):
        if inn_checksum_ok(digits):
            return digits, "", None
        return "", "", f"ИНН «{digits}» не проходит контрольную сумму"
    return "", "", None


def _classify_emails(raw: str) -> tuple[tuple[str, ...], str | None, bool]:
    """Return valid addresses, an error, and whether the cell is a phone list.

    A cell without «@» that contains digits is extra phones (the base sheet
    sometimes puts numbers in the email column). «not-an-email» stays an error.
    """
    text = raw.strip()
    if not text or text in {"—", "-", "–"}:
        return (), None, False
    if "@" not in text:
        if any(character.isdigit() for character in text):
            return (), None, True
        return (), f"некорректная эл. почта «{text}»", False
    found: list[str] = []
    seen: set[str] = set()
    for token in re.split(r"[;,\s]+", text):
        candidate = token.strip().lower()
        if not candidate:
            continue
        try:
            _EMAIL(candidate)
        except ValidationError:
            continue
        if candidate not in seen:
            seen.add(candidate)
            found.append(candidate)
    if not found:
        return (), f"некорректная эл. почта «{text}»", False
    return tuple(found), None, False


def _blank_mark(value: str) -> str:
    text = value.strip()
    if text in {"—", "-", "–"}:
        return ""
    return text


def _client_phone(phone: str) -> str:
    """First number for the contact card. Several numbers exceed the 20-digit index."""
    first = re.split(r"[;,]", phone, maxsplit=1)[0].strip()
    if len(first) <= _MAX_PHONE:
        return first
    return first[:_MAX_PHONE]


class _CompanyIndex:
    """Exact and legal-form-insensitive lookup for one import."""

    def __init__(self) -> None:
        self.by_key: dict[str, Company] = {}
        self.by_loose: dict[str, Company] = {}
        for org in Company.objects.all():
            self.add(org)

    def add(self, org: Company) -> None:
        self.by_key[org.company_key] = org
        loose = loose_company_key(org.name)
        if loose:
            self.by_loose.setdefault(loose, org)

    def find(self, name: str) -> Company | None:
        key = normalize_company_label(name)
        if key in self.by_key:
            return self.by_key[key]
        loose = loose_company_key(name)
        if len(loose) < 3:
            return None
        return self.by_loose.get(loose)


def _is_base_sheet(header: tuple[Any, ...]) -> bool:
    """True for the niche/relevance workbook, which is not the client template."""
    tokens = {header_token(cell) for cell in header}
    return "релевантность" in tokens or "нишаhoocon" in tokens


def _map_known(header_row: tuple[Any, ...], fields: dict[str, str]) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        field_name = fields.get(header_token(cell))
        if field_name and field_name not in found:
            found[field_name] = idx
    return found


def _usable_contact(name: str) -> str:
    if name.strip().casefold() in _PLACEHOLDER_CONTACTS:
        return ""
    return name.strip()


def _read_rows(file_obj: BinaryIO | BytesIO) -> tuple[list[_ClientRow], ClientImportReport]:
    """Parse the first sheet. Structural problems raise; bad data rows stay in the report."""
    from django.utils import timezone

    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises many types
        raise ClientImportError("Не удалось прочитать файл — ожидается .xlsx.") from exc

    report = ClientImportReport()
    today = timezone.localdate()
    try:
        sheet = workbook.active
        if sheet is None:
            raise ClientImportError("В книге нет листа с данными.")
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header = tuple(next(rows_iter) or ())
        except StopIteration as exc:
            raise ClientImportError("Файл пуст.") from exc
        base = _is_base_sheet(header)
        if base:
            columns = _map_known(header, _BASE_HEADER_FIELDS)
            if "company" not in columns:
                raise ClientImportError("В первой строке нужна колонка «Компания».")
        else:
            columns = _map_headers(header)
        parsed: list[_ClientRow] = []
        data_rows = 0
        for offset, raw in enumerate(rows_iter, start=2):
            cells = tuple(raw or ())
            values = {name: _cell_text(cells[idx] if idx < len(cells) else None) for name, idx in columns.items()}
            if not any(values.values()):
                report.skipped_blank += 1
                continue
            data_rows += 1
            if data_rows > MAX_CLIENT_IMPORT_ROWS:
                raise ClientImportError(
                    f"В файле больше {MAX_CLIENT_IMPORT_ROWS} строк с данными. Разбейте файл.",
                )
            row = _row_from_base(offset, values, report) if base else _row_from_template(offset, values, report, today)
            if row is None:
                continue
            problem = _row_problem(row)
            if problem:
                report.bad_rows += 1
                report.errors.append(f"Строка {row.row_number}: {problem}.")
                continue
            parsed.append(row)
        return parsed, report
    finally:
        workbook.close()


def _row_from_template(
    offset: int,
    values: dict[str, str],
    report: ClientImportReport,
    today: date,
) -> _ClientRow | None:
    """One template row. A bad INN with an email rejects the row; without an email it warns."""
    emails, email_error, email_is_phone = _classify_emails(values.get("email", ""))
    phone = values.get("phone", "")
    phone_emails, _, _ = _classify_emails(phone)
    if email_is_phone and phone_emails:
        emails = phone_emails
        phone = values.get("email", "")
        email_error = None
    elif email_is_phone:
        phone = _append_phone(phone, values.get("email", ""))
        report.warnings.append(f"Строка {offset}: в колонке почты указан телефон, карточка клиента не создана.")
    if email_error:
        report.bad_rows += 1
        report.errors.append(f"Строка {offset}: {email_error}.")
        return None
    inn, ogrn, inn_error = _classify_inn(values.get("inn", ""))
    if inn_error and emails:
        report.bad_rows += 1
        report.errors.append(f"Строка {offset}: {inn_error}.")
        return None
    if inn_error:
        report.warnings.append(f"Строка {offset}: {inn_error}. Карточка компании создана без ИНН.")
    if ogrn:
        report.warnings.append(f"Строка {offset}: в колонке ИНН указан ОГРН {ogrn}.")
    notes = values.get("notes", "")
    name = values.get("name", "").strip() or "уточнить"
    return _ClientRow(
        row_number=offset,
        company=values.get("company", ""),
        inn=inn,
        legal_address=values.get("legal_address", ""),
        name=name,
        email=emails[0] if emails else "",
        emails=emails,
        phone=phone,
        notes=notes,
        ogrn=ogrn,
        label_codes=tuple(sorted(labels_for_note(notes))),
        next_contact_on=planned_call_date(notes, today=today),
    )


def _row_from_base(offset: int, values: dict[str, str], report: ClientImportReport) -> _ClientRow | None:
    """One row of the niche workbook: whatever columns are filled, on a company card."""
    emails, email_error, email_is_phone = _classify_emails(values.get("email", ""))
    phone = _blank_mark(values.get("phone", ""))
    if email_is_phone:
        phone = _append_phone(phone, values.get("email", ""))
    if email_error:
        report.bad_rows += 1
        report.errors.append(f"Строка {offset}: {email_error}.")
        return None
    profile = values.get("profile", "").strip()
    niche = niche_title(values.get("niche", ""))
    region = _blank_mark(values.get("region", ""))
    relevance_raw = values.get("relevance", "").strip()
    lines = [
        line
        for line in (
            profile,
            f"Ниша: {niche}" if niche else "",
            f"Регион: {region}" if region else "",
            f"Релевантность: {relevance_raw}" if relevance_raw else "",
        )
        if line
    ]
    codes: list[str] = []
    niche_code = niche_label_code(niche) if niche else ""
    if niche_code:
        codes.append(niche_code)
    relevance_code = relevance_label_code(relevance_raw)
    if relevance_code:
        codes.append(relevance_code)
    region_code = region_label_code(region)
    if region_code:
        codes.append(region_code)
    contact = extract_contact_name(profile) or "уточнить"
    return _ClientRow(
        row_number=offset,
        company=values.get("company", ""),
        inn="",
        legal_address="",
        name=contact,
        email=emails[0] if emails else "",
        emails=emails,
        phone=phone,
        notes="\n".join(lines),
        website=_blank_mark(values.get("website", "")),
        city=_blank_mark(values.get("city", "")),
        niche=niche,
        region=region,
        relevance=relevance_value(relevance_code),
        label_codes=tuple(codes),
    )


def _append_phone(phone: str, extra: str) -> str:
    extra = extra.strip()
    if not extra:
        return phone
    if not phone:
        return extra
    if extra in phone:
        return phone
    return f"{phone}; {extra}"


def _apply_row(
    row: _ClientRow,
    report: ClientImportReport,
    seen_emails: set[str],
    index: _CompanyIndex,
    assignee: Any | None = None,
) -> None:
    """Insert contacts and fill the company card. An existing email is left as-is."""
    emails = row.emails
    for email in emails:
        if email in seen_emails:
            report.bad_rows += 1
            report.errors.append(f"Строка {row.row_number}: эл. почта «{email}» уже есть в этом файле.")
            return
    for email in emails:
        seen_emails.add(email)
    fresh = [email for email in emails if not Client.objects.filter(email=email).exists()]
    if emails and not fresh:
        report.skipped_existing += len(emails)
        if row.niche or row.website or row.city or row.region:
            _refresh_existing_company(row, report, index)
        return
    made = 0
    touches = 0
    created_company = False
    try:
        with transaction.atomic():
            org, created_company = _company_for_row(row, index)
            _fill_company(org, row, report)
            _claim_company(org, assignee)
            touches = apply_company_profile(org, _profile_from_row(row))
            index.add(org)
            for email in fresh:
                client = Client(
                    name=row.name or "уточнить",
                    email=email,
                    phone=_client_phone(row.phone),
                    company=org.name,
                    notes=row.notes,
                    assignee=assignee,
                )
                client.save()
                CompanyMember.objects.filter(company=org, client=client, is_confirmed=False).update(
                    is_confirmed=True,
                )
                remember_client_plan(client, row.next_contact_on)
                made += 1
    except IntegrityError:
        report.skipped_existing += 1
        return
    except DataError as exc:
        report.bad_rows += 1
        report.errors.append(f"Строка {row.row_number}: значение не помещается в карточку ({exc}).")
        return
    except _RowAbortedError:
        return
    report.created += made
    report.touches_created += touches
    if created_company:
        report.companies_created += 1
    else:
        report.companies_reused += 1


def _refresh_existing_company(row: _ClientRow, report: ClientImportReport, index: _CompanyIndex) -> None:
    """A second workbook may only add niche, site and city. Contacts stay as they are."""
    try:
        with transaction.atomic():
            org, created_company = _company_for_row(row, index)
            _fill_company(org, row, report)
            report.touches_created += apply_company_profile(org, _profile_from_row(row))
            index.add(org)
    except IntegrityError:
        return
    if created_company:
        report.companies_created += 1
    else:
        report.companies_reused += 1


def _claim_company(org: Company, assignee: Any | None) -> None:
    """A manager who uploads the row owns the company until someone else does.

    Contacts that already belong to another manager keep the company unclaimed,
    so the upload does not take over their queue.
    """
    if assignee is None or org.assignee_id is not None:
        return
    others = Client.objects.filter(company_ref=org).exclude(assignee_id=None)
    if others.exclude(assignee_id=assignee.pk).exists():
        return
    org.assignee = assignee


def _company_for_row(row: _ClientRow, index: _CompanyIndex) -> tuple[Company, bool]:
    """Find by full key or without «ООО», otherwise create. A longer legal name wins."""
    org = index.find(row.company)
    if org is None:
        org = Company(name=row.company.strip())
        org.save()
        index.add(org)
        return org, True
    incoming = row.company.strip()
    if len(incoming) > len(org.name) and loose_company_key(incoming) == loose_company_key(org.name):
        new_key = normalize_company_label(incoming)
        taken = Company.objects.exclude(pk=org.pk).filter(company_key=new_key).exists()
        if not taken:
            org.name = incoming
            org.save()
            index.add(org)
    return org, False


def _profile_from_row(row: _ClientRow) -> CompanyProfile:
    return CompanyProfile(
        contact_name=_usable_contact(row.name),
        phone=row.phone,
        emails="; ".join(row.emails),
        website=row.website,
        city=row.city,
        niche=row.niche,
        region=row.region,
        relevance=row.relevance,
        notes=row.notes,
        label_codes=set(row.label_codes),
        next_contact_on=row.next_contact_on,
        ogrn=row.ogrn,
    )


def _fill_company(org: Company, row: _ClientRow, report: ClientImportReport) -> None:
    """Write INN and legal address only into blank fields. A different INN is kept."""
    changed = False
    if row.inn:
        if not org.inn:
            org.inn = row.inn
            changed = True
        elif org.inn != row.inn:
            report.warnings.append(
                f"Строка {row.row_number}: ИНН компании «{org.name}» не изменён (в карточке уже {org.inn}).",
            )
    if row.legal_address and not org.legal_address:
        org.legal_address = row.legal_address
        changed = True
    if changed:
        org.save()


def import_clients_xlsx(
    file_obj: BinaryIO | BytesIO,
    *,
    assignee: Any | None = None,
) -> ClientImportReport:
    """Create client and company cards from the first worksheet.

    The client template (email optional) and the niche workbook are both
    accepted. Companies that differ only by «ООО» and quotes land on one card.
    Existing emails are not updated. ``assignee`` is the manager who uploaded
    the file: new contacts are set to them. An empty company owner is set to
    them only when no contact on that company already belongs to someone else.

    Args:
        file_obj: Binary ``.xlsx`` stream.
        assignee: Staff user who should own the new cards. None leaves them free.

    Returns:
        Counters and per-row messages.

    Raises:
        ClientImportError: unreadable file, missing headers, or too many rows.
    """
    rows, report = _read_rows(file_obj)
    ensure_work_labels()
    seen: set[str] = set()
    index = _CompanyIndex()
    for row in rows:
        _apply_row(row, report, seen, index, assignee)
    return report


def _write_clients_header(sheet: Worksheet) -> None:
    """Bold header row, freeze and filter — same layout as the import template."""
    sheet.title = "Клиенты"
    sheet.append(list(CLIENT_SHEET_HEADERS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G1"
    for column, width in zip("ABCDEFG", _COLUMN_WIDTHS, strict=True):
        sheet.column_dimensions[column].width = width


def export_clients_xlsx(clients: Iterable[Client]) -> bytes:
    """Write client cards to the import sheet so the file can be loaded back.

    INN and email are stored as text so Excel does not turn them into numbers.
    Company requisites come from the linked company card. A queryset is ordered
    by company name, then email; the caller decides which rows are included.

    Args:
        clients: Cards to export. Pass the admin queryset to keep manager scope.

    Returns:
        ``.xlsx`` bytes. The first sheet is «Клиенты»; there is no example sheet.
    """
    if isinstance(clients, QuerySet):
        clients = clients.select_related("company_ref").order_by("company", "email", "pk")
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    _write_clients_header(sheet)
    for index, client in enumerate(clients, start=2):
        org = client.company_ref
        sheet.append(
            [
                client.company,
                org.inn if org is not None else "",
                org.legal_address if org is not None else "",
                client.name,
                client.email,
                client.phone,
                client.notes,
            ],
        )
        sheet.cell(index, 2).number_format = "@"
        sheet.cell(index, 5).number_format = "@"
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()


def build_client_template_xlsx() -> bytes:
    """Workbook: empty «Клиенты» sheet plus an example that is not imported.

    The importer reads only the first sheet, so downloading and uploading the
    template does not create the sample company.
    """
    book = Workbook()
    data = book.active
    assert data is not None
    _write_clients_header(data)

    sample = book.create_sheet("Пример")
    sample.append(list(CLIENT_SHEET_HEADERS))
    for cell in sample[1]:
        cell.font = Font(bold=True)
    sample.append(
        [
            "ООО «Пример»",
            "7707083893",
            "123456, г. Москва, ул. Примерная, д. 1",
            "Иван Петров",
            "ivan.petrov@example.com",
            "+7 495 000-00-00",
            "Этот лист не загружается",
        ],
    )
    sample["A4"] = (
        "Этот лист не загружается. Скопируйте строки на лист «Клиенты»: "
        "одна строка — один контакт, общее название компании — одна карточка юрлица."
    )
    buf = BytesIO()
    book.save(buf)
    return buf.getvalue()
